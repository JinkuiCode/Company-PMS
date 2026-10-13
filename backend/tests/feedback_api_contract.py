"""Export workbook and API integration, using only an isolated fixture."""
import asyncio
from io import BytesIO
from pathlib import Path
import sys
from starlette.requests import Request
from openpyxl import load_workbook
from feedback_service_contract import FeedbackServiceContract


class FeedbackApiContract(FeedbackServiceContract):
    def test_export_contains_main_and_followups_without_column_filter(self):
        from app.api.feedback import export
        self.create()
        response = export(Request({'type':'http', 'method':'GET', 'path':'/api/feedback/export', 'headers':[]}), filters=None, mine=False, status_id=None, db=self.db, context=self.owner)
        async def read():
            return b''.join([part async for part in response.body_iterator])
        workbook = load_workbook(BytesIO(asyncio.run(read())), read_only=True)
        self.assertEqual(workbook.sheetnames, ['问题提报', '跟进记录'])
        self.assertEqual(list(workbook['问题提报'].values)[1][1], '列表异常')
        self.assertEqual(len(list(workbook['跟进记录'].values)), 2)
        self.assertIn('处理结论', list(workbook['问题提报'].values)[0])
        self.assertIn('处理归因', list(workbook['跟进记录'].values)[0])
        workbook.close()

    def test_http_routes_validate_inputs_and_keep_author_scope(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from app.api.feedback import router
        from app.core.database import get_db
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: self.db
        context = {'current':self.owner}
        for route in router.routes:
            for dependency in route.dependant.dependencies:
                if dependency.name == 'context':
                    app.dependency_overrides[dependency.call] = lambda: context['current']
        with TestClient(app) as client:
            self.assertEqual(client.post('/api/feedback', json={}).status_code, 422)
            response = client.post('/api/feedback', json=self.payload.model_dump())
            self.assertEqual(response.status_code, 200)
            report_id = response.json()['id']
            self.assertEqual(client.get('/api/feedback?page_size=501').status_code, 422)
            self.assertEqual(client.get('/api/feedback?page_size=500').json()['total'], 1)
            context['current'] = self.other
            self.assertEqual(client.get(f'/api/feedback/{report_id}').status_code, 404)
            self.assertEqual(client.get('/api/feedback').json()['total'], 0)


if __name__ == '__main__':
    import unittest
    unittest.main()
