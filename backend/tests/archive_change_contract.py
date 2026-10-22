"""Archive approvals against a temporary database; no live ERP writes."""
import os
from pathlib import Path
import sys
import tempfile
import unittest
import json
from uuid import uuid4
from unittest.mock import patch
from datetime import datetime, timedelta

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
TEMP = tempfile.TemporaryDirectory()
os.environ['DB_DIALECT'] = 'sqlite'
os.environ['SQLITE_DB_PATH'] = str(Path(TEMP.name) / 'unused.db')
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from app.core.database import Base, get_db
import app.models.init_db
from app.api import projects
from app.models.user import SysUser
from app.models.rbac import SysDept
from app.models.project import PmsProjectArchive, PmsProject
from app.models.product_line import SysProductLine
from app.models.erp_task import ErpSyncTask
from app.services.enum_registry import initialize_enum_definitions
from app.services.authorization import get_current_user_context


class ArchiveChangeContract(unittest.TestCase):
    def setUp(self):
        from app.api import archive_changes
        from app.models.archive_change import ArchiveChange
        self.Change = ArchiveChange
        self.engine = create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db=Session(self.engine)
        initialize_enum_definitions(self.db)
        self.dept=SysDept(dept_name='项目部',status=1)
        self.other_dept=SysDept(dept_name='其他部门',status=1)
        self.db.add_all([self.dept,self.other_dept]);self.db.flush()
        self.owner=SysUser(username='owner',real_name='项目负责人',password_hash='fixture',dept_id=self.dept.id,status=1)
        self.head=SysUser(username='head',real_name='部门负责人',password_hash='fixture',dept_id=self.other_dept.id,status=1)
        self.other=SysUser(username='other',real_name='其他用户',password_hash='fixture',dept_id=self.other_dept.id,status=1)
        self.db.add_all([self.owner,self.head,self.other]);self.db.flush()
        self.dept.leader_id=self.head.id
        self.db.add(SysProductLine(id=1,source_key='kingdee',organization_id=100,organization_code='001',organization_name='组织甲',display_name='产品线甲',name_key='line-one'))
        from app.services.archive_code_unique import lock_codes
        lock_codes(self.db.connection())
        self.archive=PmsProjectArchive(project_code='APPROVAL-001',project_name='正式名称',manager_id=self.owner.id,business_product_line_id=1,archive_category=1,machine_model=1,quantity=1,quantity_unit=7,equipment_series=1,serial_no='SN-001',created_at=datetime.now()-timedelta(days=400),erp_sync_policy='auto',is_enabled=1)
        self.db.add(self.archive);self.db.flush()
        self.linked=PmsProject(archive_id=self.archive.id,project_code=self.archive.project_code,project_name=self.archive.project_name,dept_id=self.dept.id,pm_id=self.owner.id)
        self.db.add(self.linked);self.db.commit()
        self.ctx={}
        self.actor(self.owner)
        app=FastAPI();app.include_router(archive_changes.router);app.include_router(projects.router)
        app.dependency_overrides[get_db]=lambda:self.db
        app.dependency_overrides[get_current_user_context]=lambda:self.ctx
        self.client=TestClient(app)
        self.source=patch('app.services.product_line_source.get_organization');self.source.start()

    def tearDown(self):
        self.source.stop();self.client.close();self.db.close();self.engine.dispose()

    def actor(self,user,permissions=None,lines=None):
        self.ctx.clear();self.ctx.update(user_id=user.id,dept_id=user.dept_id,data_scope=4,
            product_line_ids=[1] if lines is None else lines,
            permissions=permissions if permissions is not None else ['project:archive:view','project:archive:edit','project:archive:add'])

    def snapshot(self):
        response=self.client.get('/api/projects/archives/list',params={'archive_id':self.archive.id,'enabled':'all'})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()['items'][0]

    def payload(self,**updates):
        return {'updates':updates,'expected_version':self.snapshot()['change_version'],'reason':'业务字段变更','submission_id':str(uuid4())}

    def submit(self,**updates):
        response=self.client.post(f'/api/projects/archive-changes/{self.archive.id}/submit',json=self.payload(**updates))
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def test_owner_pending_preserves_formal_linked_values_and_queue(self):
        result=self.submit(project_name='待审批名称')
        self.assertEqual(result['status'],2)
        self.db.refresh(self.archive);self.db.refresh(self.linked)
        self.assertEqual(self.archive.project_name,'正式名称')
        self.assertEqual(self.linked.project_name,'正式名称')
        self.assertEqual(self.db.query(ErpSyncTask).count(),0)
    def test_new_archive_creation_uses_the_database_clock(self):
        from sqlalchemy import func,select
        from app.models.dict import SysDict,SysDictItem
        for code,value in [('archive_machine_model','1'),('archive_quantity_unit','7')]:
            definition=self.db.query(SysDict).filter_by(dict_code=code).one()
            self.db.add(SysDictItem(dict_id=definition.id,item_value=value,item_label='测试值',status=1))
        self.db.commit()
        response=self.client.post('/api/projects/archives',json={
            'project_code':'CLOCK-NEW','project_name':'新建时钟','product_line_id':1,'manager_id':self.owner.id,
            'archive_category':2,'machine_model':1,'quantity':1,'quantity_unit':7,
            'contract_signed_date':'2026-10-01','contract_ship_date':'2026-12-01'})
        self.assertEqual(response.status_code,200,response.text)
        created=self.db.get(PmsProjectArchive,response.json()['id']).created_at
        database_now=self.db.scalar(select(func.now()))
        self.assertLess(abs((created-database_now).total_seconds()),2)
    def test_add_permission_is_required_but_creating_is_not_limited_to_the_manager(self):
        from app.models.dict import SysDict,SysDictItem
        for code,value in [('archive_machine_model','1'),('archive_quantity_unit','7')]:
            definition=self.db.query(SysDict).filter_by(dict_code=code).one()
            self.db.add(SysDictItem(dict_id=definition.id,item_value=value,item_label='测试值',status=1))
        self.db.commit()
        payload={'project_code':'ADD-BY-OTHER','project_name':'有新增权限的其他用户',
            'product_line_id':1,'manager_id':self.owner.id,'archive_category':2,'machine_model':1,
            'quantity':1,'quantity_unit':7,'contract_signed_date':'2026-10-01','contract_ship_date':'2026-12-01'}
        self.actor(self.other,permissions=['project:archive:view'])
        self.assertEqual(self.client.post('/api/projects/archives',json=payload).status_code,403)
        self.actor(self.other,permissions=['project:archive:add'])
        response=self.client.post('/api/projects/archives',json=payload)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.db.get(PmsProjectArchive,response.json()['id']).manager_id,self.owner.id)

    def test_original_department_head_approval_atomically_applies_and_is_idempotent(self):
        result=self.submit(project_name='已审批名称',manager_id=self.other.id)
        self.assertEqual(result['approver_id'],self.head.id)
        self.actor(self.head)
        decision={'decision':'approve','expected_version':result['version'],'opinion':'同意'}
        url=f"/api/projects/archive-changes/requests/{result['id']}/decide"
        for _ in range(2):
            response=self.client.post(url,json=decision)
            self.assertEqual(response.status_code,200,response.text)
        self.db.refresh(self.archive);self.db.refresh(self.linked)
        self.assertEqual(self.archive.project_name,'已审批名称')
        self.assertEqual(self.linked.project_name,'已审批名称')
        self.assertEqual(self.archive.manager_id,self.other.id)
        self.assertEqual(self.db.query(ErpSyncTask).count(),1)

    def test_department_head_direct_save_and_other_actor_forbidden(self):
        self.actor(self.other)
        response=self.client.put(f'/api/projects/archives/{self.archive.id}',json={'customer':'不应修改','expected_version':self.snapshot()['change_version']})
        self.assertEqual(response.status_code,403,response.text)
        self.actor(self.head)
        response=self.client.put(f'/api/projects/archives/{self.archive.id}',json={'customer':'直接修改','expected_version':self.snapshot()['change_version']})
        self.assertEqual(response.status_code,200,response.text)
        self.db.refresh(self.archive)
        self.assertEqual(self.archive.customer,'直接修改')
        self.assertEqual(self.db.query(self.Change).count(),0)

    def test_same_submission_retry_and_second_active_request(self):
        payload=self.payload(project_name='待审批',manager_id=self.owner.id)
        url=f'/api/projects/archive-changes/{self.archive.id}/submit'
        first=self.client.post(url,json=payload);self.assertEqual(first.status_code,200,first.text)
        second=self.client.post(url,json=payload);self.assertEqual(second.status_code,200,second.text)
        self.assertEqual(first.json()['id'],second.json()['id'])
        third=self.client.post(url,json={**payload,'submission_id':str(uuid4())})
        self.assertEqual(third.status_code,409,third.text)
        self.assertEqual(self.db.query(self.Change).count(),1)

    def test_missing_route_and_application_failure_preserve_everything(self):
        payload=self.payload(project_name='新名称')
        self.dept.leader_id=None;self.db.commit()
        response=self.client.post(f'/api/projects/archive-changes/{self.archive.id}/submit',json=payload)
        self.assertEqual(response.status_code,422,response.text)
        self.dept.leader_id=self.head.id;self.db.commit()
        result=self.submit(project_name='新名称');self.actor(self.head)
        with patch('app.services.erp_queue.enqueue',side_effect=RuntimeError('transaction failure')):
            with self.assertRaisesRegex(RuntimeError,'transaction failure'):
                self.client.post(f"/api/projects/archive-changes/requests/{result['id']}/decide",json={'decision':'approve','expected_version':1})
        self.db.refresh(self.archive);self.db.refresh(self.linked)
        self.assertEqual(self.archive.project_name,'正式名称');self.assertEqual(self.linked.project_name,'正式名称')
        self.assertEqual(self.db.get(self.Change,result['id']).status,2)

    def test_stale_formal_version_rejects_approval_without_side_effect(self):
        result=self.submit(project_name='陈旧名称')
        self.archive.customer='其他业务路径已更新';self.db.commit()
        self.actor(self.head)
        response=self.client.post(f"/api/projects/archive-changes/requests/{result['id']}/decide",json={'decision':'approve','expected_version':1})
        self.assertEqual(response.status_code,409,response.text)
        self.assertEqual(self.archive.project_name,'正式名称')
        self.assertEqual(self.db.query(ErpSyncTask).count(),0)
        self.assertEqual(self.db.get(self.Change,result['id']).status,2)

    def test_revoke_relation_permission_scope_and_missing_route_fail_closed(self):
        result=self.submit(project_name='待审批')
        self.dept.leader_id=self.other.id;self.db.commit();self.actor(self.head)
        url=f"/api/projects/archive-changes/requests/{result['id']}/decide"
        response=self.client.post(url,json={'decision':'approve','expected_version':1})
        self.assertEqual(response.status_code,403,response.text)
        self.actor(self.owner,permissions=['project:archive:view'])
        response=self.client.post(f'/api/projects/archive-changes/{self.archive.id}/submit',json=self.payload(customer='新客户'))
        self.assertEqual(response.status_code,403,response.text)
        self.actor(self.owner,lines=[])
        response=self.client.get('/api/projects/archive-changes')
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()['total'],0)

    def test_withdraw_and_reject_keep_records_and_allow_resubmit(self):
        first=self.submit(customer='第一版')
        response=self.client.post(f"/api/projects/archive-changes/requests/{first['id']}/withdraw",json={'expected_version':1})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.db.get(self.Change,first['id']).status,5)
        second=self.submit(customer='第二版');self.actor(self.head)
        url=f"/api/projects/archive-changes/requests/{second['id']}/decide"
        response=self.client.post(url,json={'decision':'reject','expected_version':1,'opinion':'请补充'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.db.get(self.Change,second['id']).status,4)
        self.db.refresh(self.archive);self.assertIsNone(self.archive.customer)
        self.assertEqual(self.db.query(self.Change).count(),2)

    def test_current_relation_can_edit_metadata_without_admin_bypass(self):
        self.assertTrue(self.snapshot()['can_edit'])
        self.actor(self.other,permissions=['project:archive:view','project:archive:edit','business:data:all'])
        self.assertFalse(self.snapshot()['can_edit'])
    def test_completed_retries_survive_erp_pending_without_new_business_write(self):
        payload=self.payload(project_name='待审批')
        submit_url=f'/api/projects/archive-changes/{self.archive.id}/submit'
        first=self.client.post(submit_url,json=payload).json()
        self.actor(self.head)
        url=f"/api/projects/archive-changes/requests/{first['id']}/decide"
        decision={'decision':'approve','expected_version':1}
        self.assertEqual(self.client.post(url,json=decision).status_code,200)
        self.archive.erp_sync_status='pending';self.db.commit()
        self.assertEqual(self.client.post(url,json=decision).status_code,200)
        self.actor(self.owner)
        self.assertEqual(self.client.post(submit_url,json=payload).status_code,200)
        self.assertEqual(self.db.query(ErpSyncTask).count(),1)
    def test_preview_reference_differences_show_people_not_only_storage_ids(self):
        payload={'updates':{'manager_id':self.other.id},'expected_version':self.snapshot()['change_version']}
        response=self.client.post(f'/api/projects/archive-changes/{self.archive.id}/preview',json=payload)
        self.assertEqual(response.status_code,200,response.text)
        change=response.json()['differences'][0]
        self.assertEqual(change['before_display'],'项目负责人（owner）')
        self.assertEqual(change['after_display'],'其他用户（other）')
    def test_concurrent_submissions_create_only_one_pending_request(self):
        import sqlite3
        import threading
        from fastapi import HTTPException
        from app.schemas.archive_change import ArchiveChangeSubmit
        from app.services.archive_change import submit_archive_change
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'concurrent.db'
            source=self.engine.raw_connection()
            try:
                with sqlite3.connect(path) as target:source.driver_connection.backup(target)
            finally:source.close()
            engine=create_engine(f'sqlite:///{path}',connect_args={'check_same_thread':False,'timeout':5})
            barrier=threading.Barrier(2);results=[]
            payload=self.payload(customer='并发申请');scope=dict(self.ctx);actor=self.owner.id;archive_id=self.archive.id
            def send():
                with Session(engine) as db:
                    try:
                        barrier.wait(timeout=5)
                        data=ArchiveChangeSubmit(**{**payload,'submission_id':str(uuid4())})
                        submit_archive_change(db,archive_id,data,actor,scope_context=scope)
                        results.append(200)
                    except HTTPException as error:results.append(error.status_code)
                    except Exception as error:results.append(type(error).__name__)
            threads=[threading.Thread(target=send) for _ in range(2)]
            for thread in threads:thread.start()
            for thread in threads:thread.join(timeout=10)
            self.assertFalse(any(thread.is_alive() for thread in threads))
            self.assertCountEqual(results,[200,409])
            with Session(engine) as db:
                self.assertEqual(db.query(self.Change).filter_by(status=2).count(),1)
                self.assertIsNone(db.get(PmsProjectArchive,archive_id).customer)
            engine.dispose()

    def test_approval_list_query_filters_before_paging_and_rejects_unknown_field(self):
        self.submit(customer='待处理')
        filters=[{'schemaVersion':2,'field':'project_code','operator':'equals','value':'不存在'}]
        response=self.client.get('/api/projects/archive-changes',params={'filters':json.dumps(filters),'page_size':500})
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()['total'],0)
        filters[0]['field']='password'
        response=self.client.get('/api/projects/archive-changes',params={'filters':json.dumps(filters)})
        self.assertEqual(response.status_code,422,response.text)

    def test_export_checks_permissions_scope_and_visible_columns(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.submit(customer='待处理')
        payload={'columns':['project_code','status','reason']}
        url='/api/projects/archive-changes/export'
        self.assertEqual(self.client.post(url,json=payload).status_code,403)
        self.actor(self.owner,permissions=['project:archive:view','project:archive:export'])
        response=self.client.post(url,json=payload)
        self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
        sheet=load_workbook(BytesIO(response.content),read_only=True).active
        self.assertEqual(list(sheet.values),[('项目编号','审批状态','变更原因'),('APPROVAL-001','待审批','业务字段变更')])
        self.actor(self.other,permissions=['project:archive:view','project:archive:export','business:data:all'])
        response=self.client.post(url,json=payload)
        self.assertEqual(response.status_code,200)
        self.assertEqual(len(list(load_workbook(BytesIO(response.content),read_only=True).active.values)),1)
        self.assertEqual(self.client.post(url,json={'columns':['request_hash']}).status_code,422)


if __name__=='__main__':unittest.main()
